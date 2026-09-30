"""Validate ABS 2025 ERP against the project's SA4 2026 geography."""
import json
from pathlib import Path

import openpyxl
import pandas as pd

from src.acquisition.population import POPULATION_URL, CORRESPONDENCE_URL


def validate_population_mapping(population, correspondence, expected_codes):
    """Accept only complete, unique, weight-one correspondence for this release."""
    codes = set(map(str, expected_codes))
    mapping = correspondence[correspondence.SA4_CODE_2026.isin(codes)].copy()
    if set(mapping.SA4_CODE_2026) != codes:
        raise ValueError('Population correspondence has missing SA4 destinations')
    if mapping.SA4_CODE_2026.duplicated().any() or mapping.SA4_CODE_2021.duplicated().any():
        raise ValueError('Population correspondence is not one-to-one; redistribution requires a reviewed method')
    if not pd.to_numeric(mapping.RATIO_FROM_TO, errors='coerce').eq(1).all():
        raise ValueError('Population correspondence is not weight-one')
    for column in ['INDIV_TO_REGION_QLTY_INDICATOR', 'OVERALL_QUALITY_INDICATOR']:
        if not mapping[column].eq('Good').all():
            raise ValueError('Population correspondence quality is not Good')
    if not pd.to_numeric(mapping.BMOS_NULL_FLAG, errors='coerce').eq(0).all():
        raise ValueError('Population correspondence contains null geography')
    if population.sa4_code_2021.duplicated().any():
        raise ValueError('Duplicate population SA4 key')
    result = mapping.merge(population, left_on='SA4_CODE_2021', right_on='sa4_code_2021', how='left', validate='one_to_one')
    values = pd.to_numeric(result.resident_population, errors='coerce')
    if values.isna().any() or (values <= 0).any() or (values % 1 != 0).any():
        raise ValueError('Population must contain a positive integer for every spatial SA4')
    result['resident_population'] = values.astype('int64')
    return result


def build_population_silver(workbook_path, correspondence_path, sa4_path, output_path):
    regions = pd.read_parquet(sa4_path, columns=['SA4_CODE26', 'SA4_NAME26'])
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook['Table 3']
        year_row = next(sheet.iter_rows(min_row=5, max_row=5, max_col=40, values_only=True))
        headers = next(sheet.iter_rows(min_row=6, max_row=6, max_col=40, values_only=True))
        year_col = year_row.index(2025)
        code_col, name_col = headers.index('SA4 code'), headers.index('SA4 name')
        records = []
        for row in sheet.iter_rows(min_row=7, max_row=250, max_col=40, values_only=True):
            if row[0] == 1 and isinstance(row[code_col], (int, float)):
                records.append({'sa4_code_2021': str(int(row[code_col])), 'population_sa4_name': row[name_col], 'resident_population': row[year_col]})
        population = pd.DataFrame(records)
        state_sheet = workbook['Table 5']
        state_years = next(state_sheet.iter_rows(min_row=5, max_row=5, max_col=40, values_only=True))
        state_rows = list(state_sheet.iter_rows(min_row=7, max_row=20, max_col=40, values_only=True))
        state_total = next(row[state_years.index(2025)] for row in state_rows if row[0] == 1)
    finally:
        workbook.close()
    correspondence = pd.read_csv(correspondence_path, dtype={'SA4_CODE_2021': str, 'SA4_CODE_2026': str})
    result = validate_population_mapping(population, correspondence, regions.SA4_CODE26)
    if result.resident_population.sum() != state_total:
        raise ValueError('Spatial SA4 population does not reconcile to NSW ERP total')
    result = result.rename(columns={'SA4_CODE_2026': 'sa4_code', 'SA4_NAME_2026': 'sa4_name'})
    result['population_year'] = 2025
    result['population_reference_date'] = '2025-06-30'
    result['population_geography'] = 'ASGS 2021 ERP; ABS weight-one correspondence to SA4 2026'
    result = result[['sa4_code', 'sa4_name', 'resident_population', 'population_year', 'population_reference_date', 'population_geography']].sort_values('sa4_code').reset_index(drop=True)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    result.to_parquet(output_path, index=False)
    pd.testing.assert_frame_equal(result, pd.read_parquet(output_path))
    output_path.with_suffix('.metadata.json').write_text(json.dumps({
        'population_source': POPULATION_URL, 'correspondence_source': CORRESPONDENCE_URL,
        'population_reference_date': '2025-06-30', 'population_geography_year': 2021,
        'ev_geography_year': 2026, 'mapping_method': 'official_one_to_one_weight_one',
        'matched_sa4_count': len(result), 'nsw_population': int(state_total),
    }, indent=2) + '\n', encoding='utf-8')
    return result
